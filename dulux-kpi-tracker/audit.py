"""Static audit: check every cross-sheet reference lands where it is supposed to."""
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
import re, os
BASE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(BASE, "Dulux_Channel_KPI_Tracker_2026.xlsx"))
st, wk, tr, ic = wb["Setup"], wb["Weekly_Tracker"], wb["Tracking"], wb["Incentive"]
problems = []

def wk_meta(row):
    anchor = row - ((row - 6) % 3)
    return wk.cell(row=anchor, column=3).value, wk.cell(row=row, column=6).value  # KPI code, sub-type

# 1. Tracking YTD sums -> right weekly row, right sub-type, right quarter end column
QEND_EXPECT = {1: 3, 2: 6, 3: 9, 4: 12}   # months covered
kpi_of_col = {6: ("K1", "Phasing"), 7: ("K1", "Actual"), 10: ("K2", "Phasing"), 11: ("K2", "Actual"),
              15: ("K3", "Phasing"), 16: ("K3", "Actual"), 18: ("K4", "Phasing"), 19: ("K4", "Actual"),
              26: ("K5", "Phasing"), 27: ("K5", "Actual")}
pat = re.compile(r"^=SUM\(Weekly_Tracker!([A-Z]+)(\d+):([A-Z]+)(\d+)\)$")
checked = 0
for row in range(6, 66):
    q = int(str(tr.cell(row=row, column=4).value)[1])
    for col, (kexp, sexp) in kpi_of_col.items():
        fx = tr.cell(row=row, column=col).value
        m = pat.match(fx or "")
        if not m:
            problems.append(f"Tracking {get_column_letter(col)}{row}: unexpected formula {fx!r}")
            continue
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        kcode, sub = wk_meta(r1)
        if r1 != r2:
            problems.append(f"Tracking {get_column_letter(col)}{row}: row span {r1}!={r2}")
        if kcode != kexp or sub != sexp:
            problems.append(f"Tracking {get_column_letter(col)}{row}: points at {kcode}/{sub}, expected {kexp}/{sexp}")
        if c1 != "J":
            problems.append(f"Tracking {get_column_letter(col)}{row}: starts at {c1}, expected J")
        expected_end = 10 + QEND_EXPECT[q] * 5 - 1
        if column_index_from_string(c2) != expected_end:
            problems.append(f"Tracking {get_column_letter(col)}{row}: ends {c2}, expected "
                            f"{get_column_letter(expected_end)} for Q{q}")
        checked += 1
print(f"checked {checked} Tracking->Weekly sums")

# 2. Tracking person identity -> right Setup team row
for row in range(6, 66):
    p = (row - 6) // 4
    for col, letter in ((1, "C"), (2, "G"), (3, "E")):
        fx = tr.cell(row=row, column=col).value
        if f"Setup!${letter}${53 + p}" not in fx:
            problems.append(f"Tracking {get_column_letter(col)}{row}: {fx!r} not pointing at Setup row {53+p}")

# 3. Incentive -> Tracking rows (quarter order) and Setup rows
for i in range(15):
    row = 6 + i
    trow = 6 + i * 4
    for q in range(4):
        fx = ic.cell(row=row, column=9 + q).value
        if fx != f"=Tracking!$X${trow + q}":
            problems.append(f"Incentive {get_column_letter(9+q)}{row}: {fx!r} != Tracking X{trow+q}")
        qlabel = tr.cell(row=trow + q, column=4).value
        if qlabel != f"Q{q+1}":
            problems.append(f"Tracking row {trow+q} is {qlabel}, expected Q{q+1}")
    for col, exp in ((22, f"=Tracking!$U${trow+3}"), (23, f"=Tracking!$V${trow+3}")):
        if ic.cell(row=row, column=col).value != exp:
            problems.append(f"Incentive {get_column_letter(col)}{row}: {ic.cell(row=row,column=col).value!r} != {exp}")
    for col, letter in ((2, "C"), (3, "G"), (4, "E"), (5, "D"), (6, "L"), (7, "M"), (8, "N")):
        if f"Setup!${letter}${53+i}" not in (ic.cell(row=row, column=col).value or ""):
            problems.append(f"Incentive {get_column_letter(col)}{row} bad Setup ref")

# 4. Setup parameter cells carry the labels the formulas assume
expect_labels = {"C9": "Lead KPI share", "C10": "Focus KPI share", "C14": "Threshold achievement",
                 "C15": "Payout tại threshold", "C16": "Target achievement", "C17": "Payout tại target",
                 "C18": "Cap achievement", "C19": "Payout tại cap", "C22": "Gate ON / OFF",
                 "C30": "Cho phép true-up âm", "C5": "Quý đang chốt"}
for cell, label in expect_labels.items():
    lab = st[f"B{cell[1:]}"].value
    if not (lab and label.lower() in str(lab).lower()):
        problems.append(f"Setup {cell}: label is {lab!r}, expected ~{label!r}")
for i in range(5):
    if st.cell(row=44 + i, column=2).value != f"K{i+1}":
        problems.append(f"Setup E{44+i} weight row is not K{i+1}")

# 5. Weekly %Achieved rows divide the right two rows
bad = 0
for row in range(6, 231):
    if wk.cell(row=row, column=6).value != "% Achieved":
        continue
    fx = wk.cell(row=row, column=10).value
    if fx != f'=IFERROR(J{row-1}/J{row-2},"")':
        bad += 1
if bad:
    problems.append(f"{bad} %Achieved rows with unexpected formula")

print("PROBLEMS:", len(problems))
for p in problems[:30]:
    print("  -", p)
