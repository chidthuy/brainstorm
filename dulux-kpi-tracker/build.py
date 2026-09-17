# -*- coding: utf-8 -*-
"""Dulux sales-channel KPI tracker - 4 sheets: Setup, Weekly_Tracker, Tracking, Incentive."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule, CellIsRule
import os, random

BASE = os.path.dirname(os.path.abspath(__file__))
FONT = "Arial"
WHITE, NAVY, BLUE_H, GOLD_H, GREEN_H = "FFFFFF", "1F4E79", "2E75B6", "BF8F00", "548235"
LEAD_FILL, FOCUS_FILL, TRACK_FILL = "DDEBF7", "FFF2CC", "EDEDED"
ID_FILL, CALC_FILL, INPUT_FILL = "F2F2F2", "E2EFDA", "FFFDE7"
BLUE_TXT, BLACK, GREEN_TXT, RED_TXT = "0000FF", "000000", "006100", "C00000"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM, PCT, PCT0 = '#,##0;(#,##0);-', '0.0%;(0.0%);-', '0%'
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

def f(sz=10, b=False, color=BLACK, it=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=it)

def hdr(ws, row, col, text, fill=NAVY, color=WHITE, sz=9):
    c = ws.cell(row=row, column=col, value=text)
    c.font = f(sz, True, color)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BOX
    return c

def put(ws, row, col, value, *, nf=None, font=None, fill=None, align=None, border=True, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font or f()
    if nf: c.number_format = nf
    if fill: c.fill = PatternFill("solid", fgColor=fill)
    if align or wrap:
        c.alignment = Alignment(horizontal=align or "left", vertical="center", wrap_text=wrap)
    if border: c.border = BOX
    return c

def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w

def note(ws, row, col, text):
    put(ws, row, col, text, font=f(9, False, RED_TXT, it=True), border=False)

# ---------------------------------------------------------------- KPI set ----
# code, name, group, weight, unit
KPIS = [("K1", "Invoiced Value",                    "Lead",  0.50, "MVND"),
        ("K2", "Key account approach / Site check", "Lead",  0.20, "Roadshow"),
        ("K3", "Spec-in Value",                     "Focus", 0.15, "MVND"),
        ("K4", "Engaged Architect (Submit Design)", "Focus", 0.15, "D&B"),
        ("K5", "Spec-in Design (submitted design)", "Track", 0.00, "Designs")]
NK = len(KPIS)
NPEOPLE = 15

# ------------------------------------------------------------ Setup layout ---
S = "Setup"
S_CHANNEL, S_QTR, S_CCY = f"{S}!$C$4", f"{S}!$C$5", f"{S}!$C$6"
S_LEAD, S_FOCUS = f"{S}!$C$9", f"{S}!$C$10"
S_THR, S_PTHR, S_TGT, S_PTGT, S_MAX, S_PMAX = (f"{S}!$C$14", f"{S}!$C$15", f"{S}!$C$16",
                                               f"{S}!$C$17", f"{S}!$C$18", f"{S}!$C$19")
S_GATE = f"{S}!$C$22"
S_G1, S_GP1, S_G2, S_GP2, S_G3, S_GP3, S_GP4 = (f"{S}!$B$24", f"{S}!$C$24", f"{S}!$B$25",
                                                f"{S}!$C$25", f"{S}!$B$26", f"{S}!$C$26", f"{S}!$C$27")
S_CLAW = f"{S}!$C$30"
KPI_ROW0 = 44                      # K1..K5 on rows 44..48
W_KPI = [f"{S}!$E${KPI_ROW0+i}" for i in range(NK)]
TEAM_ROW0 = 53                     # 15 people on rows 53..67
TEAM_LAST = TEAM_ROW0 + NPEOPLE - 1
T_CODE, T_TEAM, T_REG, T_LM, T_NAME, T_AREA, T_STATUS = "C", "D", "E", "F", "G", "H", "I"
T_MONTHS, T_ATVP, T_PROR, T_LEADPOT, T_FOCUSPOT = "J", "K", "L", "M", "N"

def curve(ach, guard):
    return (f"IF({guard}<=0,0,IF({ach}<{S_THR},0,"
            f"IF({ach}<={S_TGT},{S_PTHR}+({ach}-{S_THR})/({S_TGT}-{S_THR})*({S_PTGT}-{S_PTHR}),"
            f"MIN({S_PMAX},{S_PTGT}+({ach}-{S_TGT})/({S_MAX}-{S_TGT})*({S_PMAX}-{S_PTGT})))))")

# --------------------------------------------------- Weekly_Tracker layout ---
WK_FIRST = 6
WCOL_FY_TGT, WCOL_TOTAL, WCOL_CHECK, WCOL_W1 = 7, 8, 9, 10      # G, H, I, J..
WCOL_LAST = WCOL_W1 + 59
def week_cols(m):
    s = WCOL_W1 + m * 5
    return s, s + 4
def wk_row(p, k, sub):
    return WK_FIRST + p * NK * 3 + k * 3 + sub
WK_LAST = wk_row(NPEOPLE - 1, NK - 1, 2)
QEND = [get_column_letter(week_cols(3 * q + 2)[1]) for q in range(4)]   # last week col of each quarter
W1L = get_column_letter(WCOL_W1)
WLL = get_column_letter(WCOL_LAST)

wb = Workbook()

# ==================================================================== SETUP ===
ws = wb.active
ws.title = "Setup"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 3, "B": 34, "C": 16, "D": 18, "E": 14, "F": 20, "G": 22, "H": 18, "I": 12,
            "J": 12, "K": 13, "L": 13, "M": 13, "N": 13, "O": 34})
ws["B2"] = "SETUP – THÔNG SỐ, KPI & NHÂN SỰ"
ws["B2"].font = f(14, True, NAVY)

def sect(r, text, span=4):
    c = put(ws, r, 2, text, font=f(10, True, WHITE), fill=BLUE_H)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=1 + span)
    c.alignment = Alignment(horizontal="left", vertical="center")

def param(r, label, value, nf=None, hint=""):
    put(ws, r, 2, label)
    put(ws, r, 3, value, nf=nf, align="center", font=f(10, True, BLUE_TXT), fill=INPUT_FILL)
    if hint:
        put(ws, r, 4, hint, font=f(9, False, "808080", it=True))

sect(3, "A. GENERAL")
param(4, "Channel", "Dulux – <tên kênh>")
param(5, "Quý đang chốt (1–4)", 3, nf="0")
param(6, "Đơn vị tiền", "MVND")

sect(8, "B. VARIABLE PAY SPLIT")
param(9, "Lead KPI share – trả theo quý", 0.70, nf=PCT0)
param(10, "Focus KPI share – trả cuối năm", 0.30, nf=PCT0)
put(ws, 11, 2, "Check", font=f(10, True))
put(ws, 11, 3, "=C9+C10", nf=PCT0, align="center", font=f(10, True), fill=CALC_FILL)
put(ws, 11, 4, '=IF(ROUND(C9+C10,6)=1,"OK","LỖI: tổng phải = 100%")', font=f(9, True, RED_TXT))

sect(13, "C. PAYOUT CURVE")
param(14, "Threshold achievement", 0.80, nf=PCT0)
param(15, "Payout tại threshold", 0.50, nf=PCT0)
param(16, "Target achievement", 1.00, nf=PCT0)
param(17, "Payout tại target", 1.00, nf=PCT0)
param(18, "Cap achievement", 1.20, nf=PCT0)
param(19, "Payout tại cap", 1.50, nf=PCT0)

sect(21, "D. PASS GATE – Focus khoá tiền Lead")
param(22, "Gate ON / OFF", "ON")
hdr(ws, 23, 2, "Focus achievement ≥"); hdr(ws, 23, 3, "% tiền Lead được trả")
for i, (b, p_) in enumerate([(1.00, 1.00), (0.90, 0.90), (0.80, 0.80), (None, 0.00)]):
    r = 24 + i
    put(ws, r, 2, b if b is not None else "< 80%", nf=PCT0 if b else None, align="center",
        font=f(10, True, BLUE_TXT), fill=INPUT_FILL)
    put(ws, r, 3, p_, nf=PCT0, align="center", font=f(10, True, BLUE_TXT), fill=INPUT_FILL)

sect(29, "E. TRUE-UP")
param(30, "Cho phép true-up âm (thu hồi) Y/N", "N")

sect(32, "F. REGION LIST")
for i, reg in enumerate(["HA NOI", "HCMC", "NORTH", "NCENT", "SCENT", "HCMEX", "MEKONG", ""]):
    put(ws, 33 + i, 2, reg, font=f(10, True, BLUE_TXT), fill=INPUT_FILL)

sect(42, "G. KPI & WEIGHT", span=5)
for i, h in enumerate(["KPI", "KPI name", "Group", "Weight", "Unit"]):
    hdr(ws, 43, 2 + i, h)
for i, (code, name, grp, wt, unit) in enumerate(KPIS):
    r = KPI_ROW0 + i
    fill = LEAD_FILL if grp == "Lead" else (FOCUS_FILL if grp == "Focus" else TRACK_FILL)
    put(ws, r, 2, code, align="center", font=f(10, True), fill=fill)
    put(ws, r, 3, name, font=f(10, True), fill=fill)
    put(ws, r, 4, grp, align="center", fill=fill)
    put(ws, r, 5, wt, nf=PCT0, align="center", font=f(10, True, BLUE_TXT), fill=INPUT_FILL)
    put(ws, r, 6, unit, align="center", fill=fill)
r = KPI_ROW0 + NK
put(ws, r, 4, "TOTAL", font=f(10, True), align="right")
put(ws, r, 5, f"=SUM(E{KPI_ROW0}:E{KPI_ROW0+NK-1})", nf=PCT0, align="center",
    font=f(10, True), fill=CALC_FILL)
put(ws, r, 6, f'=IF(ROUND(E{r},6)=1,"OK","LỖI: tổng weight phải = 100%")', font=f(9, True, RED_TXT))
put(ws, r + 1, 4, "Lead", font=f(10, True), align="right")
put(ws, r + 1, 5, f'=SUMIF($D${KPI_ROW0}:$D${KPI_ROW0+NK-1},"Lead",$E${KPI_ROW0}:$E${KPI_ROW0+NK-1})',
    nf=PCT0, align="center", fill=CALC_FILL)
put(ws, r + 1, 6, f'=IF(ROUND(E{r+1},6)=ROUND({S_LEAD},6),"OK","LỆCH so với C9")', font=f(9, True, RED_TXT))
put(ws, r + 2, 4, "Focus", font=f(10, True), align="right")
put(ws, r + 2, 5, f'=SUMIF($D${KPI_ROW0}:$D${KPI_ROW0+NK-1},"Focus",$E${KPI_ROW0}:$E${KPI_ROW0+NK-1})',
    nf=PCT0, align="center", fill=CALC_FILL)
put(ws, r + 2, 6, f'=IF(ROUND(E{r+2},6)=ROUND({S_FOCUS},6),"OK","LỆCH so với C10")', font=f(9, True, RED_TXT))

sect(51, "H. TEAM – thứ tự dòng ở đây quyết định thứ tự khối ở các sheet khác", span=13)
for i, h in enumerate(["No", "Emp Code", "Team", "Region", "Line Manager", "Name", "Area", "Status",
                       "Eligible months", "ATVP full year", "ATVP prorated", "Lead pot", "Focus pot"]):
    hdr(ws, 52, 2 + i, h)
ws.row_dimensions[52].height = 30
for i in range(NPEOPLE):
    r = TEAM_ROW0 + i
    put(ws, r, 2, f'=IF($C{r}="","",ROW()-{TEAM_ROW0-1})', align="center", nf="0")
    for c in range(3, 10):
        put(ws, r, c, None, font=f(10, False, BLUE_TXT), fill=INPUT_FILL)
    put(ws, r, 10, 12, nf="0", align="center", font=f(10, False, BLUE_TXT), fill=INPUT_FILL)
    put(ws, r, 11, None, nf=NUM, font=f(10, False, BLUE_TXT), fill=INPUT_FILL)
    put(ws, r, 12, f'=IF($C{r}="","",K{r}*MIN(J{r},12)/12)', nf=NUM)
    put(ws, r, 13, f'=IF($C{r}="","",L{r}*{S_LEAD})', nf=NUM)
    put(ws, r, 14, f'=IF($C{r}="","",L{r}*{S_FOCUS})', nf=NUM)
tr = TEAM_LAST + 1
put(ws, tr, 10, "TOTAL", font=f(10, True), align="right")
for c in (11, 12, 13, 14):
    L = get_column_letter(c)
    put(ws, tr, c, f"=SUM({L}{TEAM_ROW0}:{L}{TEAM_LAST})", nf=NUM, font=f(10, True), fill=CALC_FILL)
note(ws, tr + 2, 2, "Ô vàng = nhập. ATVP = quỹ thưởng mục tiêu cả năm (HR cấp).")

for cell, lst in (("C22", '"ON,OFF"'), ("C30", '"Y,N"'), ("C5", '"1,2,3,4"')):
    dv = DataValidation(type="list", formula1=lst, allow_blank=False, showErrorMessage=True,
                        errorTitle="Giá trị không hợp lệ", error="Chọn từ danh sách.")
    ws.add_data_validation(dv); dv.add(ws[cell])
dvr = DataValidation(type="list", formula1="=Setup!$B$33:$B$40", allow_blank=True,
                     showErrorMessage=True, errorTitle="Region không hợp lệ",
                     error="Chọn region có trong danh sách mục F.")
ws.add_data_validation(dvr); dvr.add(f"E{TEAM_ROW0}:E{TEAM_LAST}")
dvs = DataValidation(type="list", formula1='"Active,Vacant,Left,New joiner"', allow_blank=True)
ws.add_data_validation(dvs); dvs.add(f"I{TEAM_ROW0}:I{TEAM_LAST}")
dvm = DataValidation(type="whole", operator="between", formula1=0, formula2=12, allow_blank=True,
                     showErrorMessage=True, errorTitle="Số tháng không hợp lệ", error="Nhập 0–12.")
ws.add_data_validation(dvm); dvm.add(f"J{TEAM_ROW0}:J{TEAM_LAST}")
dvp = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True,
                     showErrorMessage=True, errorTitle="Giá trị âm", error="ATVP phải >= 0.")
ws.add_data_validation(dvp); dvp.add(f"K{TEAM_ROW0}:K{TEAM_LAST}")
dvw = DataValidation(type="decimal", operator="between", formula1=0, formula2=1, allow_blank=False,
                     showErrorMessage=True, errorTitle="Weight không hợp lệ", error="Weight 0–100%.")
ws.add_data_validation(dvw); dvw.add(f"E{KPI_ROW0}:E{KPI_ROW0+NK-1}")
dvc = DataValidation(type="decimal", operator="between", formula1=0, formula2=3, allow_blank=False,
                     showErrorMessage=True, errorTitle="Ngoài khoảng", error="Nhập 0–300%.")
ws.add_data_validation(dvc); dvc.add("C14:C19"); dvc.add("B24:C27"); dvc.add("C9:C10")

# =========================================================== WEEKLY TRACKER ===
ws = wb.create_sheet("Weekly_Tracker")
ws.sheet_view.showGridLines = False
ws["A2"] = "WEEKLY TRACKER – PHASING / ACTUAL / % ACHIEVED"
ws["A2"].font = f(14, True, NAVY)
note(ws, 3, 1, "Sheet nhập liệu duy nhất. Gõ FY Target (cột G) + phasing theo tuần, hàng tuần gõ Actual. "
               "Tháng 4 tuần thì để trống W5.")
widths(ws, {"A": 11, "B": 16, "C": 6, "D": 26, "E": 10, "F": 12, "G": 12, "H": 12, "I": 14})
for m in range(12):
    s, e = week_cols(m)
    for c in range(s, e + 1):
        ws.column_dimensions[get_column_letter(c)].width = 7.5
    ws.merge_cells(start_row=4, start_column=s, end_row=4, end_column=e)
    c = ws.cell(row=4, column=s, value=MONTHS[m])
    c.font = f(10, True, WHITE)
    c.fill = PatternFill("solid", fgColor=GREEN_H if m % 2 == 0 else "70AD47")
    c.alignment = Alignment(horizontal="center", vertical="center")
    for w in range(5):
        hdr(ws, 5, s + w, f"W{w+1}", fill="A9D08E", color="1F3864", sz=8)
ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=WCOL_CHECK)
c = ws.cell(row=4, column=1, value="ID / TARGET")
c.font = f(10, True, WHITE); c.fill = PatternFill("solid", fgColor="404040")
c.alignment = Alignment(horizontal="center", vertical="center")
for i, h in enumerate(["Emp Code", "Name", "KPI", "KPI name", "Unit", "Phasing / Actual",
                       "FY Target", "FY total", "Check"]):
    hdr(ws, 5, i + 1, h)
ws.row_dimensions[5].height = 26

SUBS = ["Phasing", "Actual", "% Achieved"]
for p in range(NPEOPLE):
    r0 = wk_row(p, 0, 0)
    put(ws, r0, 1, f'=IF(Setup!${T_CODE}${TEAM_ROW0+p}="","",Setup!${T_CODE}${TEAM_ROW0+p})',
        font=f(10, True, GREEN_TXT), align="center", fill="FFE699")
    put(ws, r0, 2, f'=IF(Setup!${T_NAME}${TEAM_ROW0+p}="","",Setup!${T_NAME}${TEAM_ROW0+p})',
        font=f(11, True, GREEN_TXT), fill="FFE699")
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0 + NK * 3 - 1, end_column=1)
    ws.merge_cells(start_row=r0, start_column=2, end_row=r0 + NK * 3 - 1, end_column=2)
    for k, (code, name, grp, wt, unit) in enumerate(KPIS):
        kfill = LEAD_FILL if grp == "Lead" else (FOCUS_FILL if grp == "Focus" else TRACK_FILL)
        rk = wk_row(p, k, 0)
        put(ws, rk, 3, code, align="center", font=f(10, True), fill=kfill)
        put(ws, rk, 4, name, font=f(10, True), fill=kfill, wrap=True)
        put(ws, rk, 5, unit, align="center", fill=kfill)
        for col in (3, 4, 5):
            ws.merge_cells(start_row=rk, start_column=col, end_row=rk + 2, end_column=col)
        put(ws, rk, WCOL_FY_TGT, None, nf=NUM, font=f(10, True, BLUE_TXT), fill=INPUT_FILL)
        ws.merge_cells(start_row=rk, start_column=WCOL_FY_TGT, end_row=rk + 2, end_column=WCOL_FY_TGT)
        band = ID_FILL if p % 2 == 1 else None
        for s_i, label in enumerate(SUBS):
            r = rk + s_i
            put(ws, r, 6, label, font=f(9, True), align="center",
                fill=CALC_FILL if s_i == 2 else band)
            if s_i < 2:
                put(ws, r, WCOL_TOTAL, f"=SUM({W1L}{r}:{WLL}{r})", nf=NUM, font=f(10, True), fill=CALC_FILL)
                for cidx in range(WCOL_W1, WCOL_LAST + 1):
                    put(ws, r, cidx, None, nf=NUM, font=f(10, False, BLUE_TXT), fill=INPUT_FILL)
            else:
                put(ws, r, WCOL_TOTAL, f'=IFERROR(H{r-1}/H{r-2},"")', nf=PCT, font=f(10, True), fill=CALC_FILL)
                for cidx in range(WCOL_W1, WCOL_LAST + 1):
                    L = get_column_letter(cidx)
                    put(ws, r, cidx, f'=IFERROR({L}{r-1}/{L}{r-2},"")', nf=PCT, font=f(9), fill=CALC_FILL)
        # validation column: phasing vs FY target / actual vs FY target
        anchor = wk_row(p, 0, 0)          # Emp Code is merged on the person's first row
        put(ws, rk, WCOL_CHECK,
            f'=IF($A${anchor}="","",IF($G{rk}=0,"chưa nhập FY target",'
            f'IF(ROUND(H{rk}-$G{rk},2)=0,"OK","Lệch "&TEXT(H{rk}-$G{rk},"#,##0"))))',
            font=f(9, True, RED_TXT), align="center")
        put(ws, rk + 1, WCOL_CHECK, f'=IFERROR(H{rk+1}/$G{rk},"")', nf=PCT, align="center",
            font=f(9, False, "808080"))
        put(ws, rk + 2, WCOL_CHECK, None, fill=band)
ws.freeze_panes = f"{W1L}6"

rng = f"{W1L}{WK_FIRST}:{WLL}{WK_LAST}"
for op, color, fill in [("<0.8", "9C0006", "FFC7CE"), ("<1", "9C6500", "FFEB9C"),
                        (">=1", "006100", "C6EFCE")]:
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND($F{WK_FIRST}="% Achieved",{W1L}{WK_FIRST}<>"",{W1L}{WK_FIRST}{op})'],
        fill=PatternFill("solid", fgColor=fill), font=Font(name=FONT, size=9, color=color)))
ws.conditional_formatting.add(f"I{WK_FIRST}:I{WK_LAST}", CellIsRule(
    operator="equal", formula=['"OK"'], font=Font(name=FONT, size=9, bold=True, color="006100"),
    fill=PatternFill("solid", fgColor="C6EFCE")))
dvin = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True,
                      showErrorMessage=True, errorTitle="Giá trị không hợp lệ",
                      error="Chỉ nhập số >= 0 (để trống nếu chưa có số).")
ws.add_data_validation(dvin)
dvin.add(rng)
dvin.add(f"G{WK_FIRST}:G{WK_LAST}")

# ================================================================= TRACKING ===
ws = wb.create_sheet("Tracking")
ws.sheet_view.showGridLines = False
ws["B2"] = "TRACKING – LUỸ KẾ THEO QUÝ / QUARTERLY CUMULATIVE"
ws["B2"].font = f(14, True, NAVY)
note(ws, 3, 2, "Tự động. Mỗi người 4 dòng Q1–Q4, số liệu luỹ kế từ đầu năm đến hết quý đó (Q3 = Jan–Sep).")
groups = [(1, 5, "ID", "404040"), (6, 14, "LEAD KPI – trả theo quý", BLUE_H),
          (15, 22, "FOCUS KPI – trả cuối năm", GOLD_H), (23, 25, "KẾT QUẢ", GREEN_H),
          (26, 28, "THEO DÕI – KHÔNG TÍNH ĐIỂM", "7F7F7F")]
for c1, c2, txt, hcol in groups:
    ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
    c = ws.cell(row=4, column=c1, value=txt)
    c.font = f(10, True, WHITE); c.fill = PatternFill("solid", fgColor=hcol)
    c.alignment = Alignment(horizontal="center", vertical="center")
theads = ["Emp Code", "Name", "Region", "Quarter", "Q#",
          "Invoiced Target YTD", "Invoiced Actual YTD", "Invoiced Ach %", "Invoiced factor",
          "Site check Target YTD", "Site check Actual YTD", "Site check Ach %", "Site check factor",
          "LEAD payout factor",
          "Spec-in Target YTD", "Spec-in Actual YTD", "Spec-in Ach %",
          "Engaged Target YTD", "Engaged Actual YTD", "Engaged Ach %",
          "FOCUS Ach (weighted)", "FOCUS payout factor",
          "Pass gate %", "NET LEAD factor", "Status",
          "Design Target YTD", "Design Actual YTD", "Design Ach %"]
for i, h in enumerate(theads):
    col = i + 1
    fill = NAVY
    if 6 <= col <= 14: fill = BLUE_H
    elif 15 <= col <= 22: fill = GOLD_H
    elif 23 <= col <= 25: fill = GREEN_H
    elif col >= 26: fill = "7F7F7F"
    hdr(ws, 5, col, h, fill=fill)
ws.row_dimensions[5].height = 46
widths(ws, {"A": 11, "B": 16, "C": 11, "D": 8, "E": 5})
for c in range(6, 29):
    ws.column_dimensions[get_column_letter(c)].width = 12
ws.column_dimensions["Y"].width = 13

TRK_FIRST = 6
TRK_LAST = TRK_FIRST + NPEOPLE * 4 - 1
def ytd(p, k, sub, q):
    r = wk_row(p, k, sub)
    return f"SUM(Weekly_Tracker!{W1L}{r}:{QEND[q]}{r})"

for p in range(NPEOPLE):
    for q in range(4):
        r = TRK_FIRST + p * 4 + q
        band = ID_FILL if p % 2 == 1 else None
        sr = TEAM_ROW0 + p
        put(ws, r, 1, f'=IF(Setup!${T_CODE}${sr}="","",Setup!${T_CODE}${sr})',
            font=f(10, False, GREEN_TXT), align="center", fill=band)
        put(ws, r, 2, f'=IF(Setup!${T_NAME}${sr}="","",Setup!${T_NAME}${sr})',
            font=f(10, False, GREEN_TXT), fill=band)
        put(ws, r, 3, f'=IF(Setup!${T_REG}${sr}="","",Setup!${T_REG}${sr})',
            font=f(10, False, GREEN_TXT), align="center", fill=band)
        put(ws, r, 4, f"Q{q+1}", align="center", font=f(10, True), fill=band)
        put(ws, r, 5, q + 1, align="center", nf="0", font=f(9, False, "A6A6A6"), fill=band)
        put(ws, r, 6, f"={ytd(p,0,0,q)}", nf=NUM, fill=LEAD_FILL)
        put(ws, r, 7, f"={ytd(p,0,1,q)}", nf=NUM, fill=LEAD_FILL)
        put(ws, r, 8, f"=IF(F{r}<=0,0,G{r}/F{r})", nf=PCT, font=f(10, True), fill=LEAD_FILL)
        put(ws, r, 9, f"={curve(f'H{r}', f'F{r}')}", nf=PCT, fill=LEAD_FILL)
        put(ws, r, 10, f"={ytd(p,1,0,q)}", nf=NUM, fill=LEAD_FILL)
        put(ws, r, 11, f"={ytd(p,1,1,q)}", nf=NUM, fill=LEAD_FILL)
        put(ws, r, 12, f"=IF(J{r}<=0,0,K{r}/J{r})", nf=PCT, font=f(10, True), fill=LEAD_FILL)
        put(ws, r, 13, f"={curve(f'L{r}', f'J{r}')}", nf=PCT, fill=LEAD_FILL)
        put(ws, r, 14, f"=IF(({W_KPI[0]}+{W_KPI[1]})=0,0,"
                       f"({W_KPI[0]}*I{r}+{W_KPI[1]}*M{r})/({W_KPI[0]}+{W_KPI[1]}))",
            nf=PCT, font=f(10, True), fill=CALC_FILL)
        put(ws, r, 15, f"={ytd(p,2,0,q)}", nf=NUM, fill=FOCUS_FILL)
        put(ws, r, 16, f"={ytd(p,2,1,q)}", nf=NUM, fill=FOCUS_FILL)
        put(ws, r, 17, f"=IF(O{r}<=0,0,P{r}/O{r})", nf=PCT, font=f(10, True), fill=FOCUS_FILL)
        put(ws, r, 18, f"={ytd(p,3,0,q)}", nf=NUM, fill=FOCUS_FILL)
        put(ws, r, 19, f"={ytd(p,3,1,q)}", nf=NUM, fill=FOCUS_FILL)
        put(ws, r, 20, f"=IF(R{r}<=0,0,S{r}/R{r})", nf=PCT, font=f(10, True), fill=FOCUS_FILL)
        put(ws, r, 21, f"=IF(({W_KPI[2]}+{W_KPI[3]})=0,0,"
                       f"({W_KPI[2]}*Q{r}+{W_KPI[3]}*T{r})/({W_KPI[2]}+{W_KPI[3]}))",
            nf=PCT, font=f(10, True), fill=FOCUS_FILL)
        put(ws, r, 22, f"=IF(({W_KPI[2]}+{W_KPI[3]})=0,0,({W_KPI[2]}*{curve(f'Q{r}', f'O{r}')}+"
                       f"{W_KPI[3]}*{curve(f'T{r}', f'R{r}')})/({W_KPI[2]}+{W_KPI[3]}))",
            nf=PCT, font=f(10, True), fill=FOCUS_FILL)
        put(ws, r, 23, f'=IF({S_GATE}="OFF",1,IF(U{r}>={S_G1},{S_GP1},IF(U{r}>={S_G2},{S_GP2},'
                       f'IF(U{r}>={S_G3},{S_GP3},{S_GP4}))))', nf=PCT0, align="center", fill=CALC_FILL)
        put(ws, r, 24, f"=N{r}*W{r}", nf=PCT, font=f(10, True), fill=CALC_FILL)
        put(ws, r, 25, f'=IF($A{r}="","",IF({q+1}>{S_QTR},"Chưa chốt",IF(F{r}<=0,"No target",'
                       f'IF(X{r}>=1,"On track",IF(X{r}>=0.8,"Watch","Behind")))))',
            align="center", font=f(10, True), fill=CALC_FILL)
        put(ws, r, 26, f"={ytd(p,4,0,q)}", nf=NUM, fill=TRACK_FILL)
        put(ws, r, 27, f"={ytd(p,4,1,q)}", nf=NUM, fill=TRACK_FILL)
        put(ws, r, 28, f"=IF(Z{r}<=0,0,AA{r}/Z{r})", nf=PCT, fill=TRACK_FILL)
ws.freeze_panes = "F6"

# channel roll-up at the reporting quarter
sr = TRK_LAST + 2
put(ws, sr, 2, "TỔNG KÊNH – quý chọn ở Setup!C5", font=f(11, True, NAVY), border=False)
for i, h in enumerate(["KPI", "Unit", "Target YTD", "Actual YTD", "Ach %", "Payout factor", "Weight"]):
    hdr(ws, sr + 1, 2 + i, h)
def sif(col):
    return f"SUMIF($E${TRK_FIRST}:$E${TRK_LAST},{S_QTR},${col}${TRK_FIRST}:${col}${TRK_LAST})"
for i, (code, name, grp, wt, unit, tc, ac) in enumerate(
        [(*KPIS[0], "F", "G"), (*KPIS[1], "J", "K"), (*KPIS[2], "O", "P"),
         (*KPIS[3], "R", "S"), (*KPIS[4], "Z", "AA")]):
    r = sr + 2 + i
    fill = LEAD_FILL if grp == "Lead" else (FOCUS_FILL if grp == "Focus" else TRACK_FILL)
    put(ws, r, 2, name, font=f(10, True), fill=fill)
    put(ws, r, 3, unit, align="center", fill=fill)
    put(ws, r, 4, f"={sif(tc)}", nf=NUM, fill=fill)
    put(ws, r, 5, f"={sif(ac)}", nf=NUM, fill=fill)
    put(ws, r, 6, f"=IF(D{r}<=0,0,E{r}/D{r})", nf=PCT, font=f(10, True), fill=fill)
    put(ws, r, 7, f"={curve(f'F{r}', f'D{r}')}", nf=PCT, fill=fill)
    put(ws, r, 8, f"={W_KPI[i]}", nf=PCT0, align="center", fill=fill)

for col in ["H", "L", "N", "Q", "T", "U", "X"]:
    ws.conditional_formatting.add(f"{col}{TRK_FIRST}:{col}{TRK_LAST}", ColorScaleRule(
        start_type="num", start_value=0.8, start_color="F8696B",
        mid_type="num", mid_value=1.0, mid_color="FFEB84",
        end_type="num", end_value=1.2, end_color="63BE7B"))
for state, fc, bg in [("Behind", "9C0006", "FFC7CE"), ("On track", "006100", "C6EFCE"),
                      ("Watch", "9C6500", "FFEB9C")]:
    ws.conditional_formatting.add(f"Y{TRK_FIRST}:Y{TRK_LAST}", CellIsRule(
        operator="equal", formula=[f'"{state}"'],
        font=Font(name=FONT, size=10, bold=True, color=fc), fill=PatternFill("solid", fgColor=bg)))

# ================================================================= INCENTIVE ==
ws = wb.create_sheet("Incentive")
ws.sheet_view.showGridLines = False
ws["B2"] = "INCENTIVE – LEAD theo quý + FOCUS cuối năm"
ws["B2"].font = f(14, True, NAVY)
note(ws, 3, 2, "Tự động. Tiền quý = số hưởng luỹ kế đến quý đó − số đã trả các quý trước.")
groups = [(2, 8, "ID & QUỸ THƯỞNG", "404040"), (9, 12, "NET LEAD FACTOR (luỹ kế)", BLUE_H),
          (13, 16, "HƯỞNG LUỸ KẾ", BLUE_H), (17, 21, "TIỀN TRẢ TỪNG QUÝ", NAVY),
          (22, 24, "FOCUS – CUỐI NĂM", GOLD_H), (25, 27, "TỔNG FY", GREEN_H)]
for c1, c2, txt, hcol in groups:
    ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
    c = ws.cell(row=4, column=c1, value=txt)
    c.font = f(10, True, WHITE); c.fill = PatternFill("solid", fgColor=hcol)
    c.alignment = Alignment(horizontal="center", vertical="center")
heads = ["Emp Code", "Name", "Region", "Team", "ATVP (prorated)", "Lead pot", "Focus pot",
         "Q1", "Q2", "Q3", "Q4", "Earned YTD Q1", "Earned YTD Q2", "Earned YTD Q3", "Earned YTD Q4",
         "Pay Q1", "Pay Q2", "Pay Q3", "Pay Q4", "Total Lead (đến quý chốt)",
         "Focus Ach FY", "Focus factor", "Focus payout", "TOTAL payout (đến quý chốt)", "% of ATVP", "Note"]
for i, h in enumerate(heads):
    col = i + 2
    fill = NAVY
    if 9 <= col <= 16: fill = BLUE_H
    elif 22 <= col <= 24: fill = GOLD_H
    elif col >= 25: fill = GREEN_H
    hdr(ws, 5, col, h, fill=fill)
ws.row_dimensions[5].height = 34
widths(ws, {"A": 3, "B": 11, "C": 16, "D": 11, "E": 11})
for c in range(6, 28):
    ws.column_dimensions[get_column_letter(c)].width = 12
ws.column_dimensions["AB"].width = 30

IC_FIRST = 6
IC_LAST = IC_FIRST + NPEOPLE - 1
for p in range(NPEOPLE):
    r = IC_FIRST + p
    sr = TEAM_ROW0 + p
    trow = TRK_FIRST + p * 4
    put(ws, r, 2, f'=IF(Setup!${T_CODE}${sr}="","",Setup!${T_CODE}${sr})',
        font=f(10, False, GREEN_TXT), align="center")
    put(ws, r, 3, f'=IF(Setup!${T_NAME}${sr}="","",Setup!${T_NAME}${sr})', font=f(10, False, GREEN_TXT))
    put(ws, r, 4, f'=IF(Setup!${T_REG}${sr}="","",Setup!${T_REG}${sr})',
        font=f(10, False, GREEN_TXT), align="center")
    put(ws, r, 5, f'=IF(Setup!${T_TEAM}${sr}="","",Setup!${T_TEAM}${sr})',
        font=f(10, False, GREEN_TXT), align="center")
    put(ws, r, 6, f'=IF($B{r}="",0,Setup!${T_PROR}${sr})', nf=NUM, font=f(10, False, GREEN_TXT))
    put(ws, r, 7, f'=IF($B{r}="",0,Setup!${T_LEADPOT}${sr})', nf=NUM, font=f(10, False, GREEN_TXT))
    put(ws, r, 8, f'=IF($B{r}="",0,Setup!${T_FOCUSPOT}${sr})', nf=NUM, font=f(10, False, GREEN_TXT))
    for q in range(4):
        put(ws, r, 9 + q, f"=Tracking!$X${trow + q}", nf=PCT, fill=LEAD_FILL)
        ec = get_column_letter(9 + q)
        put(ws, r, 13 + q, f"=IF({q+1}>{S_QTR},0,$G{r}*{q+1}/4*{ec}{r})", nf=NUM, fill=LEAD_FILL)
    put(ws, r, 17, f'=IF({S_CLAW}="Y",M{r},MAX(0,M{r}))', nf=NUM, font=f(10, True))
    for q in range(2, 5):
        prev, cur = get_column_letter(11 + q), get_column_letter(12 + q)
        put(ws, r, 16 + q, f'=IF({S_CLAW}="Y",{cur}{r}-{prev}{r},MAX(0,{cur}{r}-{prev}{r}))',
            nf=NUM, font=f(10, True))
    put(ws, r, 21, f"=SUM(Q{r}:T{r})", nf=NUM, font=f(10, True), fill=CALC_FILL)
    put(ws, r, 22, f"=Tracking!$U${trow + 3}", nf=PCT, fill=FOCUS_FILL)
    put(ws, r, 23, f"=Tracking!$V${trow + 3}", nf=PCT, fill=FOCUS_FILL)
    put(ws, r, 24, f"=IF({S_QTR}<4,0,$H{r}*W{r})", nf=NUM, font=f(10, True), fill=FOCUS_FILL)
    put(ws, r, 25, f"=U{r}+X{r}", nf=NUM, font=f(11, True), fill=CALC_FILL)
    put(ws, r, 26, f"=IF($F{r}<=0,0,Y{r}/$F{r})", nf=PCT, font=f(10, True), fill=CALC_FILL)
    put(ws, r, 27, f'=IF($B{r}="","",IF(MAX(I{r}:L{r})>={S_PMAX},"Chạm trần payout",'
                   f'IF(MAX(I{r}:L{r})=0,"Chưa có dữ liệu","")))', font=f(9, False, RED_TXT), wrap=True)
tot = IC_LAST + 1
put(ws, tot, 5, "TOTAL", font=f(10, True), align="right")
for c in [6, 7, 8, 13, 14, 15, 16, 17, 18, 19, 20, 21, 24, 25]:
    L = get_column_letter(c)
    put(ws, tot, c, f"=SUM({L}{IC_FIRST}:{L}{IC_LAST})", nf=NUM, font=f(10, True), fill=CALC_FILL)
ws.freeze_panes = "F6"

# =============================================================== SAMPLE DATA ==
people = [("SE-001", "SPEC", "HA NOI", "Nguyen Thi Lan", "HUNG", "Ha Noi 1", 120),
          ("SE-002", "SPEC", "HA NOI", "Nguyen Thi Lan", "DUY", "Ha Noi 2", 120),
          ("SE-003", "SPEC", "HCMC", "Tran Quang Minh", "QUYEN", "HCMC 1", 130),
          ("SE-004", "SPEC", "HCMC", "Tran Quang Minh", "Pham Minh Chau", "HCMC 2", 130),
          ("SE-005", "SPEC", "NORTH", "Nguyen Thi Lan", "Vo Thi Em", "Bac Ninh", 110),
          ("SE-006", "SPEC", "MEKONG", "Tran Quang Minh", "Ngo Van Phuc", "Can Tho", 110)]
st = wb["Setup"]
for i, (code, team, reg, lm, name, area, atvp) in enumerate(people):
    r = TEAM_ROW0 + i
    for j, v in enumerate([code, team, reg, lm, name, area]):
        st.cell(row=r, column=3 + j).value = v
    st.cell(row=r, column=9).value = "Active"
    st.cell(row=r, column=11).value = atvp

fy_targets = [[12000, 215, 4250, 80, 85], [10500, 216, 3900, 75, 78], [13800, 230, 4600, 88, 92],
              [11200, 205, 4100, 78, 82], [8600, 190, 3400, 65, 68], [7900, 180, 3100, 60, 62]]
perf = [1.03, 0.94, 1.08, 0.88, 0.99, 1.12]
season = [0.85, 0.62, 0.95, 1.00, 1.02, 1.05, 1.00, 1.03, 1.08, 1.12, 1.15, 1.13]
ssum = sum(season)
random.seed(11)
wk = wb["Weekly_Tracker"]
for p, targets in enumerate(fy_targets):
    for k, fy in enumerate(targets):
        rp, ra = wk_row(p, k, 0), wk_row(p, k, 1)
        wk.cell(row=rp, column=WCOL_FY_TGT).value = fy
        monthly = [int(round(fy * season[m] / ssum)) for m in range(12)]
        monthly[11] += fy - sum(monthly)
        for m in range(12):
            s, _ = week_cols(m)
            base = monthly[m] // 4
            weeks = [base, base, base, monthly[m] - 3 * base]
            for w in range(4):
                wk.cell(row=rp, column=s + w).value = weeks[w]
                if m < 9:
                    wk.cell(row=ra, column=s + w).value = max(
                        0, int(round(weeks[w] * perf[p] * random.uniform(0.82, 1.18))))

for name in wb.sheetnames:
    s = wb[name]
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToWidth = 1
    s.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
out = os.path.join(BASE, "Dulux_Channel_KPI_Tracker_2026.xlsx")
wb.save(out)
print("built ->", out, "| sheets:", wb.sheetnames)
