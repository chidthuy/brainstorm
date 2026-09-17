# -*- coding: utf-8 -*-
"""Part 2: Weekly_Tracker (input), Target_FY, Actual_Monthly, Tracking."""
from common import *
from openpyxl.formatting.rule import FormulaRule
import os
BASE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(BASE, "stage1.xlsx"))

SUBS = [("Phasing", INPUT_FILL, NUM), ("Actual", INPUT_FILL, NUM), ("% Achieved", CALC_FILL, PCT)]
MST_CODE, MST_NAME, MST_REG = "C", "G", "E"

# ============================================================ WEEKLY TRACKER ===
ws = wb.create_sheet("Weekly_Tracker")
ws.sheet_view.showGridLines = False
title(ws, "B2", "WEEKLY TRACKER – PHASING / ACTUAL / % ACHIEVED (sheet nhập liệu chính)")
put(ws, 3, 1, "Mỗi người 5 KPI × 3 dòng. Gõ Phasing cả năm 1 lần, hàng tuần gõ Actual. "
              "Tháng nào chỉ có 4 tuần thì để trống cột W5. Mọi sheet khác tự gom từ đây.",
    font=f(10, False, RED_TXT, it=True), border=False)
widths(ws, {"A": 11, "B": 18, "C": 7, "D": 26, "E": 10, "F": 13, "G": 12})
for m in range(12):
    s, e = week_cols(m)
    for c in range(s, e + 1):
        ws.column_dimensions[get_column_letter(c)].width = 7.5
    ws.merge_cells(start_row=4, start_column=s, end_row=4, end_column=e)
    cell = ws.cell(row=4, column=s, value=MONTHS[m])
    cell.font = f(10, True, WHITE)
    cell.fill = PatternFill("solid", fgColor=GREEN_H if m % 2 == 0 else "70AD47")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for w in range(5):
        hdr(ws, 5, s + w, f"W{w+1}", fill="A9D08E", color="1F3864", sz=8)
for i, h in enumerate(["Emp Code", "Name", "KPI", "KPI name", "Unit", "Phasing / Actual", "FY total"]):
    hdr(ws, 5, i + 1, h)
ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=7)
c = ws.cell(row=4, column=1, value="THÔNG TIN / ID")
c.font = f(10, True, WHITE); c.fill = PatternFill("solid", fgColor="404040")
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[5].height = 26

LASTW = W_FIRSTWEEK + 59
LW = get_column_letter(LASTW)
for p in range(NPEOPLE):
    r0 = wk_row(p, 0, 0)
    pband = ID_FILL if p % 2 == 1 else None
    put(ws, r0, W_CODE,
        f'=IFERROR(INDEX(Team_Master!${MST_CODE}${DATA_FIRST}:${MST_CODE}${MST_LAST},{p+1}),"")',
        font=f(10, True, GREEN_TXT), align="center", fill="FFE699")
    put(ws, r0, W_NAME,
        f'=IFERROR(INDEX(Team_Master!${MST_NAME}${DATA_FIRST}:${MST_NAME}${MST_LAST},{p+1}),"")',
        font=f(11, True, GREEN_TXT), fill="FFE699")
    ws.merge_cells(start_row=r0, start_column=W_CODE, end_row=r0 + NK * 3 - 1, end_column=W_CODE)
    ws.merge_cells(start_row=r0, start_column=W_NAME, end_row=r0 + NK * 3 - 1, end_column=W_NAME)
    for k, (code, en, vn, grp, wt, unit, src, defn) in enumerate(KPIS):
        kfill = LEAD_FILL if grp == "Lead" else (FOCUS_FILL if grp == "Focus" else TRACK_FILL)
        rk = wk_row(p, k, 0)
        put(ws, rk, W_KPI, code, align="center", font=f(10, True), fill=kfill)
        put(ws, rk, W_KPIN, en, font=f(10, True), fill=kfill, wrap=True)
        put(ws, rk, W_UNIT, unit, align="center", fill=kfill)
        for col in (W_KPI, W_KPIN, W_UNIT):
            ws.merge_cells(start_row=rk, start_column=col, end_row=rk + 2, end_column=col)
        for s_i, (label, fill, nf) in enumerate(SUBS):
            r = rk + s_i
            put(ws, r, W_TYPE, label, font=f(9, True), align="center",
                fill=CALC_FILL if s_i == 2 else pband)
            if s_i < 2:
                put(ws, r, W_FY, f"=SUM(H{r}:{LW}{r})", nf=nf, font=f(10, True), fill=CALC_FILL)
                for cidx in range(W_FIRSTWEEK, LASTW + 1):
                    put(ws, r, cidx, None, nf=nf, font=f(10, False, BLUE_TXT), fill=INPUT_FILL)
            else:
                put(ws, r, W_FY, f"=IFERROR(G{r-1}/G{r-2},\"\")", nf=PCT, font=f(10, True), fill=CALC_FILL)
                for cidx in range(W_FIRSTWEEK, LASTW + 1):
                    L = get_column_letter(cidx)
                    put(ws, r, cidx, f'=IFERROR({L}{r-1}/{L}{r-2},"")', nf=PCT,
                        font=f(9), fill=CALC_FILL)
ws.freeze_panes = "H6"
rng = f"H{DATA_FIRST}:{LW}{WK_LAST}"
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($F{DATA_FIRST}="% Achieved",H{DATA_FIRST}<>"",H{DATA_FIRST}<0.8)'],
    fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, size=9, color="9C0006")))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($F{DATA_FIRST}="% Achieved",H{DATA_FIRST}<>"",H{DATA_FIRST}<1)'],
    fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=F, size=9, color="9C6500")))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'AND($F{DATA_FIRST}="% Achieved",H{DATA_FIRST}<>"",H{DATA_FIRST}>=1)'],
    fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, size=9, color="006100")))

# ================================================ TARGET_FY / ACTUAL_MONTHLY ===
def rollup_sheet(name, ttl, sub, wsub, is_target):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    title(ws, "B2", ttl)
    put(ws, 3, 2, sub, font=f(10, False, RED_TXT if is_target else "595959", it=True), border=False)
    widths(ws, {"A": 14, "B": 11, "C": 18, "D": 7, "E": 26, "F": 9, "G": 12})
    for c in range(C_Q1, C_Q1 + 4):
        ws.column_dimensions[get_column_letter(c)].width = 11
    for c in range(C_MON1, C_MON1 + 12):
        ws.column_dimensions[get_column_letter(c)].width = 9.5
    for c in range(C_YTD1, C_YTD1 + 4):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.column_dimensions[get_column_letter(C_FYTGT)].width = 14
    ws.column_dimensions[get_column_letter(C_CHECK)].width = 18

    heads = ["Key", "Emp Code", "Name of Sales Person", "KPI", "KPI name", "Unit", "FY total",
             "Q1", "Q2", "Q3", "Q4"] + MONTHS + ["YTD Q1", "YTD Q2", "YTD Q3", "YTD Q4"]
    if is_target:
        heads += ["FY Target (nhập)", "Check phasing"]
    for i, h in enumerate(heads):
        col = i + 1
        fill = NAVY
        if C_Q1 <= col <= C_Q1 + 3:
            fill = BLUE_H
        elif C_MON1 <= col <= C_MON1 + 11:
            fill = GREEN_H
        elif C_YTD1 <= col <= C_YTD1 + 3:
            fill = "808080"
        elif col >= C_FYTGT:
            fill = GOLD_H
        hdr(ws, 5, col, h, fill=fill)
    ws.row_dimensions[5].height = 26
    put(ws, 4, C_MON1, "TỰ GOM TỪ WEEKLY_TRACKER →", font=f(9, True, "375623"), border=False)

    for p in range(NPEOPLE):
        for k, (code, en, vn, grp, wt, unit, src, defn) in enumerate(KPIS):
            r = DATA_FIRST + p * NK + k
            band = ID_FILL if p % 2 == 1 else None
            kfill = LEAD_FILL if grp == "Lead" else (FOCUS_FILL if grp == "Focus" else TRACK_FILL)
            put(ws, r, C_KEY, f'=IF($B{r}="","",$B{r}&"|"&$D{r})', font=f(9, False, "A6A6A6"), fill=band)
            put(ws, r, C_CODE,
                f'=IFERROR(INDEX(Team_Master!${MST_CODE}${DATA_FIRST}:${MST_CODE}${MST_LAST},{p+1}),"")',
                font=f(10, False, GREEN_TXT), align="center", fill=band)
            put(ws, r, C_NAME,
                f'=IFERROR(INDEX(Team_Master!${MST_NAME}${DATA_FIRST}:${MST_NAME}${MST_LAST},{p+1}),"")',
                font=f(10, False, GREEN_TXT), fill=band)
            put(ws, r, C_KPI, code, align="center", font=f(10, True), fill=kfill)
            put(ws, r, C_KPIN, en, fill=kfill, wrap=True)
            put(ws, r, C_UNIT, unit, align="center", fill=kfill)
            q1c = get_column_letter(C_Q1)
            q4c = get_column_letter(C_Q1 + 3)
            put(ws, r, C_FY, f"=SUM({q1c}{r}:{q4c}{r})", nf=NUM, font=f(10, True), fill=CALC_FILL)
            for q in range(4):
                a = get_column_letter(C_MON1 + q * 3)
                b = get_column_letter(C_MON1 + q * 3 + 2)
                put(ws, r, C_Q1 + q, f"=SUM({a}{r}:{b}{r})", nf=NUM, font=f(10, True), fill=LEAD_FILL)
            wr = wk_row(p, k, wsub)
            for m in range(12):
                s, e = week_cols(m)
                put(ws, r, C_MON1 + m,
                    f"=SUM(Weekly_Tracker!{get_column_letter(s)}{wr}:{get_column_letter(e)}{wr})",
                    nf=NUM, font=f(10, False, GREEN_TXT), fill=band)
            for q in range(4):
                prev = get_column_letter(C_YTD1 + q - 1)
                cur = get_column_letter(C_Q1 + q)
                put(ws, r, C_YTD1 + q,
                    f"={cur}{r}" if q == 0 else f"={prev}{r}+{cur}{r}", nf=NUM, fill=band)
            if is_target:
                put(ws, r, C_FYTGT, None, nf=NUM, font=f(10, True, BLUE_TXT), fill=INPUT_FILL)
                fy, tg = get_column_letter(C_FY), get_column_letter(C_FYTGT)
                put(ws, r, C_CHECK,
                    f'=IF($B{r}="","",IF({tg}{r}=0,"chưa nhập FY target",'
                    f'IF(ROUND({fy}{r}-{tg}{r},2)=0,"OK","Lệch "&TEXT({fy}{r}-{tg}{r},"#,##0"))))',
                    font=f(9, True, RED_TXT), align="center")
    ws.freeze_panes = "G6"

    tr = TGT_LAST + 3
    put(ws, tr - 1, 2, "TỔNG TOÀN KÊNH THEO KPI / CHANNEL TOTAL BY KPI", font=f(11, True, NAVY), border=False)
    for i, h in enumerate(["KPI", "KPI name", "Unit", "FY total", "Q1", "Q2", "Q3", "Q4"]):
        hdr(ws, tr, 4 + i, h)
    for i, (code, en, vn, grp, wt, unit, src, defn) in enumerate(KPIS):
        r = tr + 1 + i
        put(ws, r, 4, code, align="center", font=f(10, True))
        put(ws, r, 5, en, font=f(10, True))
        put(ws, r, 6, unit, align="center")
        for j, col in enumerate([get_column_letter(C_FY)] + [get_column_letter(C_Q1 + q) for q in range(4)]):
            put(ws, r, 7 + j,
                f'=SUMIF($D${DATA_FIRST}:$D${TGT_LAST},$D{r},{col}${DATA_FIRST}:{col}${TGT_LAST})',
                nf=NUM, font=f(10, True), fill=CALC_FILL)
    return ws

rollup_sheet("Target_FY", "TARGET / PHASING FY2026 – gom từ Weekly_Tracker",
             "Chỉ nhập cột 'FY Target (nhập)' – các cột còn lại gom từ phasing tuần. "
             "Cột 'Check phasing' báo lệch nếu phasing cả năm chưa khớp FY target.", 0, True)
rollup_sheet("Actual_Monthly", "ACTUAL FY2026 – gom từ Weekly_Tracker",
             "Sheet chỉ đọc. Actual nhập ở Weekly_Tracker (dòng 'Actual').", 1, False)

# ================================================================= TRACKING ====
W_K1, W_K2 = "KPI_Framework!$F$6", "KPI_Framework!$F$7"
W_K3, W_K4 = "KPI_Framework!$F$8", "KPI_Framework!$F$9"
ws = wb.create_sheet("Tracking")
ws.sheet_view.showGridLines = False
title(ws, "B2", "PERFORMANCE TRACKING – LUỸ KẾ THEO QUÝ / QUARTERLY CUMULATIVE")
put(ws, 3, 2, "Sheet tự động 100%. Mỗi người 4 dòng Q1–Q4; số liệu LUỸ KẾ từ đầu năm đến hết quý đó "
              "(Q3 = Jan–Sep). Target luỹ kế = phasing luỹ kế trên Weekly_Tracker.",
    font=f(10, False, "595959", it=True), border=False)
groups = [(1, 5, "THÔNG TIN / ID", "404040"), (6, 14, "LEAD KPI – 70% – trả theo quý", BLUE_H),
          (15, 22, "FOCUS KPI – 30% – trả cuối năm", GOLD_H), (23, 25, "KẾT QUẢ / RESULT", GREEN_H),
          (26, 28, "THEO DÕI – KHÔNG TÍNH ĐIỂM", "7F7F7F")]
for c1, c2, txt, hcol in groups:
    ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
    cell = ws.cell(row=4, column=c1, value=txt)
    cell.font = f(10, True, WHITE); cell.fill = PatternFill("solid", fgColor=hcol)
    cell.alignment = Alignment(horizontal="center", vertical="center")
theads = ["Emp Code", "Name", "Region", "Quarter", "Q#",
          "Invoiced Target YTD", "Invoiced Actual YTD", "Invoiced Ach %", "Invoiced factor",
          "Site check Target YTD", "Site check Actual YTD", "Site check Ach %", "Site check factor",
          "LEAD payout factor",
          "Spec-in Target YTD", "Spec-in Actual YTD", "Spec-in Ach %",
          "Engaged Target YTD", "Engaged Actual YTD", "Engaged Ach %",
          "FOCUS Ach (weighted)", "FOCUS payout factor",
          "Pass gate %", "NET LEAD factor", "Status",
          "Design Target YTD", "Design Actual YTD", "Design Ach %"]
for i, h in enumerate(theads):
    col = i + 1
    fill = NAVY
    if 6 <= col <= 14:
        fill = BLUE_H
    elif 15 <= col <= 22:
        fill = GOLD_H
    elif 23 <= col <= 25:
        fill = GREEN_H
    elif col >= 26:
        fill = "7F7F7F"
    hdr(ws, 5, col, h, fill=fill)
ws.row_dimensions[5].height = 46
widths(ws, {"A": 11, "B": 18, "C": 11, "D": 8, "E": 5})
for c in range(6, 29):
    ws.column_dimensions[get_column_letter(c)].width = 12
ws.column_dimensions["Y"].width = 13

def look(sheet, kpi, row):
    return (f'IFERROR(INDEX({sheet}!{YTD_RANGE},'
            f'MATCH($A{row}&"|{kpi}",{sheet}!$A${DATA_FIRST}:$A${TGT_LAST},0),$E{row}),0)')

for r in range(DATA_FIRST, TRK_LAST + 1):
    p = (r - DATA_FIRST) // 4
    band = ID_FILL if p % 2 == 1 else None
    put(ws, r, 1, f'=IFERROR(INDEX(Team_Master!${MST_CODE}${DATA_FIRST}:${MST_CODE}${MST_LAST},'
                  f'INT((ROW()-{DATA_FIRST})/4)+1),"")', font=f(10, False, GREEN_TXT), align="center", fill=band)
    put(ws, r, 2, f'=IFERROR(INDEX(Team_Master!${MST_NAME}${DATA_FIRST}:${MST_NAME}${MST_LAST},'
                  f'INT((ROW()-{DATA_FIRST})/4)+1),"")', font=f(10, False, GREEN_TXT), fill=band)
    put(ws, r, 3, f'=IFERROR(INDEX(Team_Master!${MST_REG}${DATA_FIRST}:${MST_REG}${MST_LAST},'
                  f'INT((ROW()-{DATA_FIRST})/4)+1),"")', font=f(10, False, GREEN_TXT), align="center", fill=band)
    put(ws, r, 4, f'="Q"&E{r}', align="center", font=f(10, True), fill=band)
    put(ws, r, 5, f"=MOD(ROW()-{DATA_FIRST},4)+1", align="center", nf="0", font=f(9, False, "A6A6A6"), fill=band)
    # Lead – K1 invoiced value
    put(ws, r, 6, f"={look('Target_FY','K1',r)}", nf=NUM, fill=LEAD_FILL)
    put(ws, r, 7, f"={look('Actual_Monthly','K1',r)}", nf=NUM, fill=LEAD_FILL)
    put(ws, r, 8, f"=IF(F{r}<=0,0,G{r}/F{r})", nf=PCT, font=f(10, True), fill=LEAD_FILL)
    put(ws, r, 9, f"={curve(f'H{r}', f'F{r}')}", nf=PCT, fill=LEAD_FILL)
    # Lead – K2 site check
    put(ws, r, 10, f"={look('Target_FY','K2',r)}", nf=NUM, fill=LEAD_FILL)
    put(ws, r, 11, f"={look('Actual_Monthly','K2',r)}", nf=NUM, fill=LEAD_FILL)
    put(ws, r, 12, f"=IF(J{r}<=0,0,K{r}/J{r})", nf=PCT, font=f(10, True), fill=LEAD_FILL)
    put(ws, r, 13, f"={curve(f'L{r}', f'J{r}')}", nf=PCT, fill=LEAD_FILL)
    put(ws, r, 14, f"=IF(({W_K1}+{W_K2})=0,0,({W_K1}*I{r}+{W_K2}*M{r})/({W_K1}+{W_K2}))",
        nf=PCT, font=f(10, True), fill=CALC_FILL)
    # Focus – K3 spec-in value, K4 engaged architect
    put(ws, r, 15, f"={look('Target_FY','K3',r)}", nf=NUM, fill=FOCUS_FILL)
    put(ws, r, 16, f"={look('Actual_Monthly','K3',r)}", nf=NUM, fill=FOCUS_FILL)
    put(ws, r, 17, f"=IF(O{r}<=0,0,P{r}/O{r})", nf=PCT, font=f(10, True), fill=FOCUS_FILL)
    put(ws, r, 18, f"={look('Target_FY','K4',r)}", nf=NUM, fill=FOCUS_FILL)
    put(ws, r, 19, f"={look('Actual_Monthly','K4',r)}", nf=NUM, fill=FOCUS_FILL)
    put(ws, r, 20, f"=IF(R{r}<=0,0,S{r}/R{r})", nf=PCT, font=f(10, True), fill=FOCUS_FILL)
    put(ws, r, 21, f"=IF(({W_K3}+{W_K4})=0,0,({W_K3}*Q{r}+{W_K4}*T{r})/({W_K3}+{W_K4}))",
        nf=PCT, font=f(10, True), fill=FOCUS_FILL)
    put(ws, r, 22, f"=IF(({W_K3}+{W_K4})=0,0,({W_K3}*{curve(f'Q{r}', f'O{r}')}+"
                   f"{W_K4}*{curve(f'T{r}', f'R{r}')})/({W_K3}+{W_K4}))",
        nf=PCT, font=f(10, True), fill=FOCUS_FILL)
    put(ws, r, 23, f'=IF({P_GATE}="OFF",1,IF(U{r}>={P_G1},{P_GP1},IF(U{r}>={P_G2},{P_GP2},'
                   f'IF(U{r}>={P_G3},{P_GP3},{P_GP4}))))', nf=PCT0, align="center", fill=CALC_FILL)
    put(ws, r, 24, f"=N{r}*W{r}", nf=PCT, font=f(10, True), fill=CALC_FILL)
    put(ws, r, 25, f'=IF($A{r}="","",IF(F{r}<=0,"No target",IF(X{r}>=1,"On track",'
                   f'IF(X{r}>=0.8,"Watch","Behind"))))', align="center", font=f(10, True), fill=CALC_FILL)
    # tracking-only K5
    put(ws, r, 26, f"={look('Target_FY','K5',r)}", nf=NUM, fill=TRACK_FILL)
    put(ws, r, 27, f"={look('Actual_Monthly','K5',r)}", nf=NUM, fill=TRACK_FILL)
    put(ws, r, 28, f"=IF(Z{r}<=0,0,AA{r}/Z{r})", nf=PCT, fill=TRACK_FILL)
ws.freeze_panes = "F6"

wb.save(os.path.join(BASE, "stage2.xlsx"))
print("part2 ok")
