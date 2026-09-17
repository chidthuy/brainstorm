# -*- coding: utf-8 -*-
"""Shared styles / constants for the Dulux channel KPI tracker build."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
WHITE = "FFFFFF"
NAVY = "1F4E79"
BLUE_H = "2E75B6"
GOLD_H = "BF8F00"
GREEN_H = "548235"
LEAD_FILL = "DDEBF7"
FOCUS_FILL = "FFF2CC"
TRACK_FILL = "EDEDED"
ID_FILL = "F2F2F2"
CALC_FILL = "E2EFDA"
INPUT_FILL = "FFFDE7"
NOTE_FILL = "FCE4D6"

BLUE_TXT = "0000FF"
BLACK = "000000"
GREEN_TXT = "006100"
RED_TXT = "C00000"

thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def f(sz=10, b=False, color=BLACK, it=False):
    return Font(name=F, size=sz, bold=b, color=color, italic=it)

def title(ws, cell, text, sz=14):
    ws[cell] = text
    ws[cell].font = f(sz, True, NAVY)

def hdr(ws, row, col, text, fill=NAVY, color=WHITE, wrap=True, sz=9):
    c = ws.cell(row=row, column=col, value=text)
    c.font = f(sz, True, color)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    c.border = BOX
    return c

def put(ws, row, col, value, *, nf=None, font=None, fill=None, align=None, border=True, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font or f()
    if nf:
        c.number_format = nf
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if align or wrap:
        c.alignment = Alignment(horizontal=align or "left", vertical="center", wrap_text=wrap)
    if border:
        c.border = BOX
    return c

def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w

def band_fill(row, col1, col2, ws, color):
    for c in range(col1, col2 + 1):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=color)

NUM = '#,##0;(#,##0);-'
PCT = '0.0%;(0.0%);-'
PCT0 = '0%'
MONEY = '#,##0;(#,##0);-'
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# ---------------- Parameter cells -------------------------------------------
P = "Parameters"
P_LEAD  = f"{P}!$C$10"
P_FOCUS = f"{P}!$C$11"
P_THR   = f"{P}!$C$15"
P_PTHR  = f"{P}!$C$16"
P_TGT   = f"{P}!$C$17"
P_PTGT  = f"{P}!$C$18"
P_MAX   = f"{P}!$C$19"
P_PMAX  = f"{P}!$C$20"
P_GATE  = f"{P}!$C$23"
P_G1, P_GP1 = f"{P}!$B$25", f"{P}!$C$25"
P_G2, P_GP2 = f"{P}!$B$26", f"{P}!$C$26"
P_G3, P_GP3 = f"{P}!$B$27", f"{P}!$C$27"
P_GP4  = f"{P}!$C$28"
P_CLAW = f"{P}!$C$31"
P_QTR  = f"{P}!$C$7"

def curve(ach, guard):
    """Payout factor from an achievement ref; 0 when the target (guard ref) is empty."""
    return (f"IF({guard}<=0,0,"
            f"IF({ach}<{P_THR},0,"
            f"IF({ach}<={P_TGT},{P_PTHR}+({ach}-{P_THR})/({P_TGT}-{P_THR})*({P_PTGT}-{P_PTHR}),"
            f"MIN({P_PMAX},{P_PTGT}+({ach}-{P_TGT})/({P_MAX}-{P_TGT})*({P_PMAX}-{P_PTGT})))))")

# ---------------- KPI set ----------------------------------------------------
# code, EN name, VN name, group, weight, unit, data source, definition
KPIS = [
    ("K1", "Invoiced Value", "Doanh số xuất hoá đơn", "Lead", 0.50, "MVND",
     "DERP / Invoice report",
     "Giá trị hoá đơn thực xuất cho khách hàng / dự án thuộc danh sách phụ trách."),
    ("K2", "Key account approach / Site check", "Site check – tiếp cận key account", "Lead", 0.20,
     "Roadshow", "Mini Sales app / CRM",
     "Số lượt roadshow / site check hợp lệ (có check-in, có biên bản trên app)."),
    ("K3", "Spec-in Value", "Giá trị dự án được chỉ định", "Focus", 0.15, "MVND", "CRM",
     "Giá trị dự án Dulux được spec-in trong năm (chuẩn ~50 MVND / design)."),
    ("K4", "Engaged Architect (Submit Design)", "Kiến trúc sư engaged (nộp design)", "Focus", 0.15,
     "D&B", "CRM",
     "Số KTS/nhà thầu D&B engaged – tính theo tài khoản duy nhất có submit design trong năm."),
    ("K5", "Spec-in Design (submitted design)", "Số design được nộp", "Tracking", 0.00, "Designs", "CRM",
     "Chỉ số dẫn dắt của K3 (design x 50 MVND). Weight 0 – chỉ theo dõi, chưa tính điểm."),
]
KCODES = [k[0] for k in KPIS]
NK = len(KPIS)

# ---------------- sheet geometry --------------------------------------------
NPEOPLE = 15
DATA_FIRST = 6
MST_LAST = DATA_FIRST + NPEOPLE - 1                 # Team_Master        6..20
TGT_LAST = DATA_FIRST + NPEOPLE * NK - 1            # Target / Actual    6..80
TRK_LAST = DATA_FIRST + NPEOPLE * 4 - 1             # Tracking           6..65
WK_LAST = DATA_FIRST + NPEOPLE * NK * 3 - 1         # Weekly_Tracker     6..230
IC_FIRST, IC_LAST = DATA_FIRST, DATA_FIRST + NPEOPLE - 1

# Target / Actual column map (identical on both sheets)
C_KEY, C_CODE, C_NAME, C_KPI, C_KPIN, C_UNIT = 1, 2, 3, 4, 5, 6
C_FY, C_Q1, C_MON1, C_YTD1 = 7, 8, 12, 24          # FY | Q1..Q4 | Jan..Dec | YTD Q1..Q4
C_FYTGT, C_CHECK = 28, 29                           # Target_FY only
YTD_RANGE = f"$X${DATA_FIRST}:$AA${TGT_LAST}"

# Weekly_Tracker column map
W_CODE, W_NAME, W_KPI, W_KPIN, W_UNIT, W_TYPE, W_FY = 1, 2, 3, 4, 5, 6, 7
W_FIRSTWEEK = 8                                     # 12 months x 5 week slots = 60 cols
def week_cols(month_idx):
    s = W_FIRSTWEEK + month_idx * 5
    return s, s + 4

def wk_row(person, kpi_idx, sub):
    """sub: 0 = Phasing, 1 = Actual, 2 = % Achieved."""
    return DATA_FIRST + person * NK * 3 + kpi_idx * 3 + sub
