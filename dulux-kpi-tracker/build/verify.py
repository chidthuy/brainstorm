from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import os
BASE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(BASE, "Dulux_Channel_KPI_Incentive_Tracker_2026_Template.xlsx"), data_only=True)
km, tm, wk, tg, ac, tr, ic, db = (wb["KPI_Framework"], wb["Team_Master"], wb["Weekly_Tracker"],
                                  wb["Target_FY"], wb["Actual_Monthly"], wb["Tracking"],
                                  wb["Incentive_Calc"], wb["Dashboard"])
print("weights:", [km.cell(row=r, column=6).value for r in range(6, 11)], "| total", km["F11"].value,
      "|", km["H11"].value, "| lead", km["F12"].value, "| focus", km["F13"].value)
print("Master SE-001: ATVP", tm["L6"].value, "prorated", tm["M6"].value,
      "lead pot", tm["N6"].value, "focus pot", tm["O6"].value)
print("Weekly SE-001 K1 phasing FY:", wk["G6"].value, "| actual FY:", wk["G7"].value,
      "| %ach FY:", wk["G8"].value, "| Jan W1 phasing:", wk["H6"].value, "actual:", wk["H7"].value,
      "%:", wk["H8"].value)
print("Target_FY r6:", "FY", tg["G6"].value, "Q1", tg["H6"].value, "Q3", tg["J6"].value,
      "YTDQ3", tg["Z6"].value, "FYtgt", tg["AB6"].value, "check", tg["AC6"].value)
print("Target_FY r7 (K2 site check): FY", tg["G7"].value, "FYtgt", tg["AB7"].value, "check", tg["AC7"].value)
print("Actual r6: FY", ac["G6"].value, "YTDQ3", ac["Z6"].value, "YTDQ4", ac["AA6"].value)
hdrs = [tr.cell(row=5, column=c).value for c in range(1, 29)]
for row in (6, 8, 9):
    vals = [tr.cell(row=row, column=c).value for c in range(1, 29)]
    print(f"Tracking r{row}:", {h: (round(v, 3) if isinstance(v, float) else v)
                                for h, v in zip(hdrs, vals) if v not in (None, 0, "")})
print("Incentive SE-001:", [ic.cell(row=6, column=c).value for c in range(2, 28)])
print("Incentive TOTAL row:", [(get_column_letter(c), ic.cell(row=21, column=c).value) for c in (6,7,8,21,24,25)])
print("Dashboard channel block:")
for r in range(7, 13):
    print("  ", [db.cell(row=r, column=c).value for c in range(2, 10)])
print("Dashboard regions:")
for r in range(16, 25):
    v = [db.cell(row=r, column=c).value for c in range(2, 14)]
    if v[0]: print("  ", v)
print("Dashboard people:")
for r in range(29, 36):
    v = [db.cell(row=r, column=c).value for c in range(2, 14)]
    if v[1]: print("  ", v)
