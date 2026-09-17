"""Recompute the whole chain in Python from the literal inputs and compare with the engine."""
import json, os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
BASE = os.path.dirname(os.path.abspath(__file__))
V = json.load(open(os.path.join(BASE, "calc_values.json")))
wb = load_workbook(os.path.join(BASE, "Dulux_Channel_KPI_Tracker_2026.xlsx"))
wk, st = wb["Weekly_Tracker"], wb["Setup"]
def g(sheet, cell):
    return V.get(f"{sheet.upper()}!{cell}")

THR, PTHR, TGT, PTGT, MX, PMAX = .80, .50, 1.00, 1.00, 1.20, 1.50
GATES = [(1.00, 1.00), (0.90, 0.90), (0.80, 0.80)]
W = {"K1": .50, "K2": .20, "K3": .15, "K4": .15, "K5": 0.0}

def curve(a, guard):
    if guard <= 0: return 0.0
    if a < THR: return 0.0
    if a <= TGT: return PTHR + (a - THR) / (TGT - THR) * (PTGT - PTHR)
    return min(PMAX, PTGT + (a - TGT) / (MX - TGT) * (PMAX - PTGT))

def wk_row(p, k, sub): return 6 + p * 15 + k * 3 + sub
def week_cols(m): return 10 + m * 5, 14 + m * 5

def ytd(p, k, sub, q):
    r = wk_row(p, k, sub)
    tot = 0
    for m in range(3 * (q + 1)):
        s, e = week_cols(m)
        for c in range(s, e + 1):
            v = wk.cell(row=r, column=c).value
            if isinstance(v, (int, float)): tot += v
    return tot

def gate(focus_ach):
    for b, pay in GATES:
        if focus_ach >= b: return pay
    return 0.0

fails, checks = [], 0
def eq(label, got, exp, tol=1e-6):
    global checks
    checks += 1
    if got is None or abs(float(got) - exp) > max(tol, abs(exp) * 1e-9):
        fails.append(f"{label}: file={got} expected={exp}")

for p in range(6):                       # the six sample people
    for q in range(4):
        row = 6 + p * 4 + q
        t1, a1 = ytd(p, 0, 0, q), ytd(p, 0, 1, q)
        t2, a2 = ytd(p, 1, 0, q), ytd(p, 1, 1, q)
        t3, a3 = ytd(p, 2, 0, q), ytd(p, 2, 1, q)
        t4, a4 = ytd(p, 3, 0, q), ytd(p, 3, 1, q)
        eq(f"Tracking F{row} invoice target YTD", g("Tracking", f"F{row}"), t1)
        eq(f"Tracking G{row} invoice actual YTD", g("Tracking", f"G{row}"), a1)
        ach1 = a1 / t1 if t1 > 0 else 0
        eq(f"Tracking H{row} invoice ach", g("Tracking", f"H{row}"), ach1)
        eq(f"Tracking I{row} invoice factor", g("Tracking", f"I{row}"), curve(ach1, t1))
        ach2 = a2 / t2 if t2 > 0 else 0
        eq(f"Tracking L{row} site ach", g("Tracking", f"L{row}"), ach2)
        lead = (W["K1"] * curve(ach1, t1) + W["K2"] * curve(ach2, t2)) / (W["K1"] + W["K2"])
        eq(f"Tracking N{row} lead factor", g("Tracking", f"N{row}"), lead)
        ach3 = a3 / t3 if t3 > 0 else 0
        ach4 = a4 / t4 if t4 > 0 else 0
        focus_ach = (W["K3"] * ach3 + W["K4"] * ach4) / (W["K3"] + W["K4"])
        focus_f = (W["K3"] * curve(ach3, t3) + W["K4"] * curve(ach4, t4)) / (W["K3"] + W["K4"])
        eq(f"Tracking U{row} focus ach", g("Tracking", f"U{row}"), focus_ach)
        eq(f"Tracking V{row} focus factor", g("Tracking", f"V{row}"), focus_f)
        gt = gate(focus_ach)
        eq(f"Tracking W{row} gate", g("Tracking", f"W{row}"), gt)
        eq(f"Tracking X{row} net lead", g("Tracking", f"X{row}"), lead * gt)

    # incentive chain
    irow = 6 + p
    atvp = st.cell(row=53 + p, column=11).value
    lead_pot, focus_pot = atvp * .70, atvp * .30
    eq(f"Incentive F{irow} ATVP", g("Incentive", f"F{irow}"), atvp)
    eq(f"Incentive G{irow} lead pot", g("Incentive", f"G{irow}"), lead_pot)
    eq(f"Incentive H{irow} focus pot", g("Incentive", f"H{irow}"), focus_pot)
    RQ = 3   # Setup!C5
    earned, paid_prev, total_pay = [], 0, 0
    for q in range(4):
        net = g("Tracking", f"X{6 + p*4 + q}")
        e = 0.0 if (q + 1) > RQ else lead_pot * (q + 1) / 4 * net
        earned.append(e)
        eq(f"Incentive {get_column_letter(13+q)}{irow} earned YTD Q{q+1}",
           g("Incentive", f"{get_column_letter(13+q)}{irow}"), e)
        pay = max(0, e - paid_prev)
        eq(f"Incentive {get_column_letter(17+q)}{irow} pay Q{q+1}",
           g("Incentive", f"{get_column_letter(17+q)}{irow}"), pay)
        paid_prev = e
        total_pay += pay
    eq(f"Incentive U{irow} total lead", g("Incentive", f"U{irow}"), total_pay)
    focus_pay = 0.0 if RQ < 4 else focus_pot * g("Tracking", f"V{6 + p*4 + 3}")
    eq(f"Incentive X{irow} focus payout", g("Incentive", f"X{irow}"), focus_pay)
    eq(f"Incentive Y{irow} total FY", g("Incentive", f"Y{irow}"), total_pay + focus_pay)
    eq(f"Incentive Z{irow} % of ATVP", g("Incentive", f"Z{irow}"), (total_pay + focus_pay) / atvp)

# phasing vs FY target validation
for p in range(6):
    for k in range(5):
        r = wk_row(p, k, 0)
        fy = wk.cell(row=r, column=7).value
        eq(f"Weekly H{r} phasing total", g("Weekly_Tracker", f"H{r}"), fy)
        chk = g("Weekly_Tracker", f"I{r}")
        checks += 1
        if chk != "OK":
            fails.append(f"Weekly I{r} check = {chk!r}, expected OK")

print(f"CHECKS RUN: {checks}   FAILURES: {len(fails)}")
for x in fails[:25]:
    print("  -", x)
print("\nSample readout – SE-001 (HUNG):")
for q in range(4):
    row = 6 + q
    print(f"  Q{q+1}: invoice {g('Tracking', f'F{row}'):,.0f} vs {g('Tracking', f'G{row}'):,.0f} "
          f"= {g('Tracking', f'H{row}'):.1%} | site {g('Tracking', f'L{row}'):.1%} | "
          f"lead factor {g('Tracking', f'N{row}'):.1%} | focus {g('Tracking', f'U{row}'):.1%} | "
          f"gate {g('Tracking', f'W{row}'):.0%} | net {g('Tracking', f'X{row}'):.1%} | "
          f"{g('Tracking', f'Y{row}')}")
print("  Incentive:", {k: round(g("Incentive", f"{c}6"), 2) if isinstance(g("Incentive", f"{c}6"), float)
                       else g("Incentive", f"{c}6")
                       for k, c in [("lead pot", "G"), ("pay Q1", "Q"), ("pay Q2", "R"), ("pay Q3", "S"),
                                    ("pay Q4", "T"), ("total lead", "U"), ("focus", "X"),
                                    ("TOTAL FY", "Y"), ("% ATVP", "Z")]})
print("\nChannel rollup at Q3 (Tracking bottom block):")
for r in range(68, 73):
    print("   ", [wb["Tracking"].cell(row=r, column=2).value,
                  g("Tracking", f"D{r}"), g("Tracking", f"E{r}"),
                  round(g("Tracking", f"F{r}"), 3) if isinstance(g("Tracking", f"F{r}"), float) else None,
                  round(g("Tracking", f"G{r}"), 3) if isinstance(g("Tracking", f"G{r}"), float) else None])
